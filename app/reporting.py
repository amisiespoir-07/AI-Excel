"""报告生成。

支持两种模式，由模型定义的 ``report`` 段决定：

* **内置版式**：``report.template`` 为空时使用。小节顺序由 ``report.sections`` 配置，
  程序不再写死「必须有哪几节」。
* **Word 模板驱动**：``report.template`` 指向一个 .docx 模板，报告内容由模板里的占位符决定。

模板占位符约定
--------------

标量占位符（可出现在正文段落或表格单元格里）::

    {{job_id}}  {{model_id}}  {{model_name}}  {{model_version}}
    {{generated_at}}  {{title}}  {{summary}}  {{note}}

表格标记：在某一行**第一格**写 ``{{parameters_table}}`` 或 ``{{results_table}}``，
该行即成为行模板并被复制填充：

* 标记所在格渲染成该行的 ``display_name``；
* 同一行其余格可按 ``{{key}}`` / ``{{display_name}}`` / ``{{value}}`` / ``{{unit}}`` 渲染。

示例（3 列表格）::

    参数                     数值         单位
    {{parameters_table}}    {{value}}    {{unit}}

填充完成后，标记行本身会被移除。未被模板引用的占位符保持原样，方便定位模板错误。
"""

from __future__ import annotations

import copy
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from .domain import ModelDefinition


_PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")
_TABLE_MARKER = re.compile(r"\{\{\s*(parameters_table|results_table)\s*\}\}")
_SUMMARY_FALLBACK = "结果说明暂不可用，请以计算结果表为准。"


def _is_windows() -> bool:
    import os

    return os.name == "nt"


# --------------------------------------------------------------------------- #
# 行数据
# --------------------------------------------------------------------------- #

def parameter_rows(model: ModelDefinition, parameters: dict[str, Any]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for key, value in parameters.items():
        spec = model.parameters.get(key)
        rows.append(
            {
                "key": key,
                "display_name": f"{spec.display_name} — {spec.description}" if spec else key,
                "value": str(value),
                "unit": (spec.unit or "") if spec else "",
            }
        )
    return rows


def result_rows(model: ModelDefinition, results: dict[str, Any]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for key, value in results.items():
        definition = model.results.get(key)
        rows.append(
            {
                "key": key,
                "display_name": definition.display_name if definition else key,
                "value": str(value),
                "unit": (definition.unit or "") if definition else "",
            }
        )
    return rows


# --------------------------------------------------------------------------- #
# 模板渲染
# --------------------------------------------------------------------------- #

def _substitute(text: str, mapping: dict[str, Any]) -> str:
    return _PLACEHOLDER.sub(
        lambda match: str(mapping.get(match.group(1), match.group(0))), text
    )


def _render_paragraph(paragraph, mapping: dict[str, Any]) -> None:
    original = paragraph.text
    if "{{" not in original:
        return
    rendered = _substitute(original, mapping)
    if rendered == original:
        return
    if paragraph.runs:
        # 保留首个 run 的格式，清空其余 run：占位符常被 Word 拆散到多个 run 里。
        paragraph.runs[0].text = rendered
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = rendered


def _render_scalars(document, mapping: dict[str, Any]) -> None:
    for paragraph in document.paragraphs:
        _render_paragraph(paragraph, mapping)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _render_paragraph(paragraph, mapping)


def _expand_marker_row(table, marker_row, rows: list[dict[str, str]]) -> None:
    """把标记行复制成 len(rows) 行，然后删掉标记行本身。"""
    from docx.table import _Row

    template_tr = copy.deepcopy(marker_row._tr)
    for values in rows:
        new_tr = copy.deepcopy(template_tr)
        marker_row._tr.addprevious(new_tr)
        for cell in _Row(new_tr, table).cells:
            text = _TABLE_MARKER.sub("{{display_name}}", cell.text)
            cell.text = _substitute(text, values)
    marker_row._tr.getparent().remove(marker_row._tr)


def _fill_marked_tables(document, rows_by_kind: dict[str, list[dict[str, str]]]) -> set[str]:
    filled: set[str] = set()
    for table in document.tables:
        for row in table.rows:
            if not row.cells:
                continue
            marker = _TABLE_MARKER.search(row.cells[0].text)
            if marker is None:
                continue
            kind = marker.group(1)
            if kind in filled:
                continue
            _expand_marker_row(table, row, rows_by_kind.get(kind, []))
            filled.add(kind)
    return filled


def _safe_report_template(root: Path, template: str) -> Path:
    root = Path(root).resolve()
    candidate = (root / template).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise RuntimeError("报告模板路径超出允许的模板目录") from exc
    if candidate.suffix.lower() != ".docx":
        raise RuntimeError("报告模板必须是 .docx 文件")
    if not candidate.is_file():
        raise RuntimeError(f"找不到报告模板：{template}")
    return candidate


# --------------------------------------------------------------------------- #
# 内置版式
# --------------------------------------------------------------------------- #

def _add_rows_table(document, headers: list[str], rows: list[dict[str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for item in rows:
        cells = table.add_row().cells
        cells[0].text = item["display_name"]
        cells[1].text = item["value"]
        cells[2].text = item["unit"]


def _build_builtin_report(document, model: ModelDefinition, rows_by_kind, summary: str) -> None:
    for section in model.report.sections:
        document.add_heading(section.title, level=1)
        if section.kind == "parameters":
            _add_rows_table(document, ["参数", "数值", "单位"], rows_by_kind["parameters_table"])
        elif section.kind == "results":
            _add_rows_table(document, ["结果字段", "结果", "单位"], rows_by_kind["results_table"])
        else:
            document.add_paragraph(summary)


# --------------------------------------------------------------------------- #
# 公共入口
# --------------------------------------------------------------------------- #

def export_word_preview(docx_path: Path, pdf_path: Path) -> Path:
    """Export a DOCX report as a PDF sidecar using Microsoft Word COM."""
    if not _is_windows():
        raise RuntimeError("Word PDF 预览导出仅支持 Windows + Microsoft Word")
    if not docx_path.is_file():
        raise FileNotFoundError(f"Word 报告不存在：{docx_path.name}")

    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("缺少 pywin32，无法调用 Microsoft Word 导出 PDF") from exc

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(docx_path.resolve()),
            ReadOnly=True,
            AddToRecentFiles=False,
            ConfirmConversions=False,
            NoEncodingDialog=True,
        )
        # wdExportFormatPDF = 17. Keep arguments minimal; optional COM None
        # values have caused E_INVALIDARG in Office automation elsewhere.
        document.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
    except Exception as exc:
        raise RuntimeError(f"Microsoft Word 导出预览 PDF 失败：{exc}") from exc
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass

    if not pdf_path.is_file() or pdf_path.stat().st_size < 5:
        raise RuntimeError("Microsoft Word 未生成有效预览 PDF")
    with pdf_path.open("rb") as file:
        if file.read(5) != b"%PDF-":
            raise RuntimeError("Microsoft Word 导出的预览文件不是有效 PDF")
    return pdf_path


def generate_report(
    output_path: Path,
    model: ModelDefinition,
    parameters: dict[str, Any],
    results: dict[str, Any],
    summary: str,
    job_id: str,
    *,
    template_root: Path | None = None,
    generated_at: datetime | None = None,
) -> Path:
    """按模型定义的 ``report`` 段生成 Word 报告，返回实际写入的路径。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，请安装 requirements.txt") from exc

    report = model.report
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    timestamp = generated_at or datetime.now().astimezone()
    mapping: dict[str, Any] = {
        "job_id": job_id,
        "model_id": model.model.id,
        "model_name": model.model.name,
        "model_version": str(model.model.version),
        "generated_at": timestamp.strftime("%Y-%m-%d %H:%M:%S %Z"),
        "title": report.title or f"{model.model.name}计算报告",
        "summary": (summary or "").strip() or _SUMMARY_FALLBACK,
        "note": report.note,
    }
    rows_by_kind = {
        "parameters_table": parameter_rows(model, parameters),
        "results_table": result_rows(model, results),
    }

    if report.template:
        if template_root is None:
            from .config import get_settings

            template_root = get_settings().report_template_root
        template_path = _safe_report_template(Path(template_root), report.template)
        document = Document(str(template_path))
        _fill_marked_tables(document, rows_by_kind)
        _render_scalars(document, mapping)
    else:
        document = Document()
        document.add_heading(mapping["title"], level=0)
        document.add_paragraph(f"任务编号：{job_id}")
        document.add_paragraph(f"模型版本：{model.model.id} / v{model.model.version}")
        document.add_paragraph(f"生成时间：{mapping['generated_at']}")
        _build_builtin_report(document, model, rows_by_kind, mapping["summary"])
        document.add_paragraph(report.note)

    document.core_properties.title = mapping["title"]
    document.save(str(output_path))
    return output_path
