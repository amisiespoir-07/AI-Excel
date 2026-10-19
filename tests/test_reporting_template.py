"""报告层解耦的回归测试：版式由配置决定，模板由文件决定。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from app.domain import ModelDefinition, load_model_definition
from app.reporting import generate_report

ROOT = Path(__file__).resolve().parents[1]


def thermal_model(**report_overrides) -> ModelDefinition:
    model = load_model_definition(ROOT / "config" / "models" / "model_energy_saving_calc" / "model.json")
    payload = model.model_dump(mode="json")
    payload["report"].update(report_overrides)
    return ModelDefinition.model_validate(payload)


PARAMETERS = {"heat_load": 240.0, "fuel_hydrogen_content": 15.3}
RESULTS = {"thermal_result": 123.4, "flue_gas_flow": 55.5}


def document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def test_default_builtin_layout_covers_parameters_results_and_summary(tmp_path):
    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(),
        PARAMETERS,
        RESULTS,
        "AI 说明文本",
        "job-1",
    )
    text = document_text(path)
    # 标题来自模型定义的 report.title，不在测试里抄写
    assert thermal_model().report.title in text
    assert "job-1" in text
    assert "一、输入参数" in text
    assert "二、计算结果" in text
    assert "三、结果说明" in text
    assert "AI 说明文本" in text
    assert "热负荷" in text and "240.0" in text
    assert "123.4" in text


def test_sections_are_configurable_not_hardcoded(tmp_path):
    sections = [
        {"kind": "results", "title": "甲、结果优先"},
        {"kind": "parameters", "title": "乙、输入在后"},
    ]
    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(sections=sections),
        PARAMETERS,
        RESULTS,
        "should-not-appear",
        "job-2",
    )
    text = document_text(path)
    assert "甲、结果优先" in text
    assert "乙、输入在后" in text
    # 没有配置 summary 小节，就不该出现说明小节，也不该出现说明正文
    assert "三、结果说明" not in text
    assert "should-not-appear" not in text
    # 顺序也要跟着配置走
    assert text.index("甲、结果优先") < text.index("乙、输入在后")


def test_title_and_note_are_configurable(tmp_path):
    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(title="自定义标题", note="自定义免责声明"),
        PARAMETERS,
        RESULTS,
        "说明",
        "job-3",
    )
    text = document_text(path)
    assert "自定义标题" in text
    assert "自定义免责声明" in text
    assert "计算报告" not in text.replace("自定义标题", "")


def test_empty_summary_falls_back_to_local_text(tmp_path):
    path = generate_report(
        tmp_path / "report.docx", thermal_model(), PARAMETERS, RESULTS, "   ", "job-4"
    )
    assert "结果说明暂不可用" in document_text(path)


def write_template(path: Path) -> None:
    document = Document()
    document.add_heading("{{title}}", level=0)
    document.add_paragraph("任务编号：{{job_id}}")
    document.add_paragraph("模型：{{model_name}} v{{model_version}}")
    document.add_paragraph("生成时间：{{generated_at}}")

    parameters_table = document.add_table(rows=1, cols=3)
    header = parameters_table.rows[0].cells
    header[0].text = "参数"
    header[1].text = "数值"
    header[2].text = "单位"
    row = parameters_table.add_row().cells
    row[0].text = "{{parameters_table}}"
    row[1].text = "{{value}}"
    row[2].text = "{{unit}}"

    results_table = document.add_table(rows=1, cols=3)
    result_header = results_table.rows[0].cells
    result_header[0].text = "结果键"
    result_header[1].text = "结果"
    result_header[2].text = "单位"
    result_row = results_table.add_row().cells
    result_row[0].text = "{{results_table}}"
    result_row[1].text = "{{key}}={{value}}"
    result_row[2].text = "{{unit}}"

    document.add_paragraph("结论：{{summary}}")
    document.save(str(path))


def test_word_template_drives_layout_via_placeholders(tmp_path):
    template_root = tmp_path / "reports"
    template_root.mkdir()
    write_template(template_root / "custom.docx")

    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(template="custom.docx", title="模板标题"),
        PARAMETERS,
        RESULTS,
        "模板说明",
        "job-9",
        template_root=template_root,
    )
    text = document_text(path)

    # 标量占位符全部落实
    assert "模板标题" in text
    assert "job-9" in text
    assert f"v{thermal_model().model.version}" in text
    assert "模板说明" in text
    assert "{{" not in text, f"模板占位符未被完全替换：\n{text}"

    # 参数表按模板行复制出两行
    assert "热负荷" in text
    assert "240.0" in text
    assert "MW" in text
    assert "氢气含量" in text and "15.3" in text

    # 结果表使用 {{key}}={{value}} 自定义格式
    assert "thermal_result=123.4" in text


def test_template_marker_row_is_removed_after_expansion(tmp_path):
    template_root = tmp_path / "reports"
    template_root.mkdir()
    write_template(template_root / "custom.docx")

    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(template="custom.docx"),
        PARAMETERS,
        RESULTS,
        "说明",
        "job-10",
        template_root=template_root,
    )
    document = Document(str(path))
    parameters_table = document.tables[0]
    # 表头 1 行 + 2 个参数行；标记行必须已被删除
    assert len(parameters_table.rows) == 3
    first_column = [row.cells[0].text for row in parameters_table.rows]
    assert first_column[0] == "参数"
    assert "{{parameters_table}}" not in first_column[1]


def test_template_without_markers_is_left_alone(tmp_path):
    template_root = tmp_path / "reports"
    template_root.mkdir()
    document = Document()
    document.add_paragraph("只有标量：{{job_id}}")
    document.save(str(template_root / "plain.docx"))

    path = generate_report(
        tmp_path / "report.docx",
        thermal_model(template="plain.docx"),
        PARAMETERS,
        RESULTS,
        "说明",
        "job-11",
        template_root=template_root,
    )
    text = document_text(path)
    assert "只有标量：job-11" in text


def test_template_path_escape_is_rejected(tmp_path):
    template_root = tmp_path / "reports"
    template_root.mkdir()
    outside = tmp_path / "outside.docx"
    Document().save(str(outside))

    with pytest.raises(RuntimeError, match="超出允许的模板目录"):
        generate_report(
            tmp_path / "report.docx",
            thermal_model(template="../outside.docx"),
            PARAMETERS,
            RESULTS,
            "说明",
            "job-12",
            template_root=template_root,
        )


def test_template_must_be_docx_and_exist(tmp_path):
    template_root = tmp_path / "reports"
    template_root.mkdir()
    (template_root / "notes.txt").write_text("not a docx", encoding="utf-8")

    with pytest.raises(RuntimeError, match="必须是 .docx"):
        generate_report(
            tmp_path / "report.docx",
            thermal_model(template="notes.txt"),
            PARAMETERS,
            RESULTS,
            "说明",
            "job-13",
            template_root=template_root,
        )

    with pytest.raises(RuntimeError, match="找不到报告模板"):
        generate_report(
            tmp_path / "report.docx",
            thermal_model(template="missing.docx"),
            PARAMETERS,
            RESULTS,
            "说明",
            "job-14",
            template_root=template_root,
        )


def test_report_defaults_do_not_change_existing_models():
    model = load_model_definition(ROOT / "config" / "models" / "model_energy_saving_calc" / "model.json")
    assert model.report.enabled is True
    assert model.report.template is None
    assert model.report.ai_summary is True
    assert model.report.pdf_preview is True
    assert model.report.file_name == "calculation-report.docx"
    assert [section.kind for section in model.report.sections] == [
        "parameters",
        "results",
        "summary",
    ]


def test_report_section_kind_is_validated():
    model = load_model_definition(ROOT / "config" / "models" / "model_energy_saving_calc" / "model.json")
    payload = deepcopy(model.model_dump(mode="json"))
    payload["report"]["sections"] = [{"kind": "not_a_kind", "title": "x"}]
    with pytest.raises(ValueError):
        ModelDefinition.model_validate(payload)
